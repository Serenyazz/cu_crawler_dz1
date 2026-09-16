import os
import re
import shutil
from datetime import datetime
from pathlib import Path
import pandas as pd
import xlrd
from openpyxl.styles import numbers

# Базовая папка — каталог, где лежит сам скрипт (transport_tracker)
BASE_DIR = Path(__file__).resolve().parent

def parse_date(date_str: str):
    """Преобразует строку вида DD.MM.YYYY в объект date."""
    if not date_str:
        return None
    try:
        return datetime.strptime(date_str.strip(), '%d.%m.%Y').date()
    except ValueError:
        return None

def parse_datetime(datetime_str: str):
    """Преобразует строку вида DD.MM.YYYY HH:MM:SS в объект datetime."""
    if not datetime_str:
        return None
    for fmt in ('%d.%m.%Y %H:%M:%S', '%d.%m.%Y %H:%M', '%d.%m.%Y'):
        try:
            return datetime.strptime(datetime_str.strip(), fmt)
        except ValueError:
            continue
    return None

def extract_trip_data(filepath: Path) -> dict:
    wb = xlrd.open_workbook(filepath)
    sheet = wb.sheet_by_index(0)
    
    extracted = {
        'Номер заявки': None,
        'Дата заявки': None,
        'Номер рейса': None,
        'Дата и время рейса': None,
        'Гос номер ТС': None,
        'ФИО водителя': None,
        'Стоимость перевозки': None,
        'Имя файла': str(filepath.name)
    }
    
    for r in range(sheet.nrows):
        row_values = [str(sheet.cell_value(r, c)).strip() for c in range(sheet.ncols)]
        row_text = " ".join([v for v in row_values if v])
        
        # 1. Номер заявки (строка с нулями) и Дата заявки (date)
        if 'Заявка №' in row_text and extracted['Номер заявки'] is None:
            match = re.search(r'Заявка\s*№\s*(\S+)\s*от\s*(\d{2}\.\d{2}\.\d{4})', row_text)
            if match:
                extracted['Номер заявки'] = str(match.group(1)).strip()
                extracted['Дата заявки'] = parse_date(match.group(2))
                
        # 2. Номер рейса (int) и Дата/время рейса (datetime)
        if 'Рейс' in row_text and extracted['Номер рейса'] is None:
            match = re.search(r'Рейс\s*([\d\s]+?)\s*от\s*(\d{2}\.\d{2}\.\d{4}(?:\s+\d{2}:\d{2}:\d{2})?)', row_text)
            if match:
                digits_only = re.sub(r'\D', '', match.group(1))
                extracted['Номер рейса'] = int(digits_only) if digits_only else None
                extracted['Дата и время рейса'] = parse_datetime(match.group(2))
                
        first_cell = row_values[0] if len(row_values) > 0 else ""
        second_cell = sheet.cell_value(r, 1) if sheet.ncols > 1 else ""
        
        # 3. Гос номер ТС (строка)
        if 'Гос номер ТС' in first_cell and extracted['Гос номер ТС'] is None:
            extracted['Гос номер ТС'] = str(second_cell).strip()
            
        # 4. ФИО водителя (строка)
        elif 'ФИО водителя' in first_cell and extracted['ФИО водителя'] is None:
            extracted['ФИО водителя'] = str(second_cell).strip()
            
        # 5. Стоимость перевозки (десятичное число float)
        elif 'Стоимость перевозки' in first_cell and extracted['Стоимость перевозки'] is None:
            try:
                extracted['Стоимость перевозки'] = round(float(second_cell), 2)
            except (ValueError, TypeError):
                cleaned_price = re.sub(r'[^\d.]', '', str(second_cell).replace(',', '.'))
                extracted['Стоимость перевозки'] = float(cleaned_price) if cleaned_price else None
                
    return extracted

def update_registry():
    inbox_path = BASE_DIR / 'inbox'
    archive_path = BASE_DIR / 'processed'
    registry_path = BASE_DIR / 'реестр_рейсов.xlsx'
    
    inbox_path.mkdir(exist_ok=True)
    archive_path.mkdir(exist_ok=True)
    
    files = [
        f for f in inbox_path.iterdir() 
        if f.is_file() and f.suffix.lower() in ('.xls', '.xlsx')
    ]
    
    if not files:
        print(f"Входящих файлов в '{inbox_path}' не обнаружено.")
        return
    
    # Чтение существующего реестра со строгим сохранением типа 'Номер заявки'
    if registry_path.exists():
        existing_df = pd.read_excel(registry_path, dtype={'Номер заявки': str})
        existing_orders = set(existing_df['Номер заявки'].dropna().astype(str).str.strip())
    else:
        existing_df = pd.DataFrame()
        existing_orders = set()
        
    new_records = []
    
    for file in files:
        try:
            data = extract_trip_data(file)
            order_no = str(data.get('Номер заявки') or '').strip()
            
            if order_no and order_no in existing_orders:
                print(f"Заявка {order_no} ({file.name}) уже учтена в реестре. Пропуск.")
            else:
                new_records.append(data)
                if order_no:
                    existing_orders.add(order_no)
                print(f"Успешно обработан: {file.name} -> Заявка № {order_no}")
                
            shutil.move(str(file), str(archive_path / file.name))
        except Exception as err:
            print(f"Ошибка при обработке {file.name}: {err}")

    if new_records:
        new_df = pd.DataFrame(new_records)
        combined_df = pd.concat([existing_df, new_df], ignore_index=True) if not existing_df.empty else new_df
        
        # Гарантируем строковый тип для номера заявки при объединении
        combined_df['Номер заявки'] = combined_df['Номер заявки'].astype(str)
        
        # Запись в Excel и форматирование колонок
        with pd.ExcelWriter(registry_path, engine='openpyxl') as writer:
            combined_df.to_excel(writer, index=False, sheet_name='Рейсы')
            ws = writer.sheets['Рейсы']
            
            # Применение числовых и датовых форматов Excel
            for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
                for cell in row:
                    col_name = ws.cell(row=1, column=cell.column).value
                    
                    if col_name == 'Номер заявки':
                        cell.number_format = '@'  # Текстовый формат, нули не срежутся
                    elif col_name == 'Дата заявки':
                        cell.number_format = 'DD.MM.YYYY'
                    elif col_name == 'Дата и время рейса':
                        cell.number_format = 'DD.MM.YYYY HH:MM:SS'
                    elif col_name == 'Номер рейса':
                        cell.number_format = '0'
                    elif col_name == 'Стоимость перевозки':
                        cell.number_format = '#,##0.00'
            
            # Автоподбор ширины колонок
            for col in ws.columns:
                max_len = max(len(str(cell.value or '')) for cell in col)
                col_letter = col[0].column_letter
                ws.column_dimensions[col_letter].width = max(max_len + 4, 14)
                
        print(f"\nРеестр обновлен: {registry_path} (добавлено {len(new_records)} записей).")
    else:
        print("\nНовых записей для внесения не обнаружено.")

if __name__ == '__main__':
    update_registry()